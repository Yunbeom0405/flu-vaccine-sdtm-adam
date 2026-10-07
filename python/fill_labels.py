"""Fill the label placeholders in the SDTM programs from specs/sdtm-spec.xlsx."""
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
ws = openpyxl.load_workbook(ROOT / "specs/sdtm-spec.xlsx")["Variables"]
labels = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    labels.setdefault(r[1], {}).setdefault(r[2], r[3])


def block(dom):
    w = max(len(n) for n in labels[dom])
    rows = [f"    {n.lower():<{w}} = '{t}'" for n, t in labels[dom].items()]
    return "  label\n" + "\n".join(rows) + ";"


for f in sorted((ROOT / "sas/sdtm").glob("*.sas")):
    s = f.read_text()
    s2 = re.sub(r"/\*LABELS (\w+)\*/", lambda m: block(m.group(1)), s)
    if s2 != s:
        f.write_text(s2)
        print(f.name, "labels filled")
