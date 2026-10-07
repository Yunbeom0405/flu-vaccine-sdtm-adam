"""Build P3 ADaM spec workbook (same layout as P2 adam-spec.xlsx)."""
from collections import OrderedDict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "projects/P3-vaccine-immuno/specs/adam-spec.xlsx"
SDTM_SPEC = ROOT / "projects/P3-vaccine-immuno/specs/sdtm-spec.xlsx"
CT = "CDISC/NCI CT 2026-09-25"
CT_OLD = "CDISC/NCI CT 2025-09-26"

datasets = [
    ("ADSL", "Subject-Level Analysis Dataset", "SUBJECT LEVEL ANALYSIS DATASET", "One record per subject",
     "STUDYID, USUBJID", "COM.ADSL.POP", "Source: sdtm dm, ds, ex, is. All 267 screened subjects."),
    ("ADIS", "Immunogenicity Analysis Dataset", "BASIC DATA STRUCTURE",
     "One record per subject per strain per visit",
     "STUDYID, USUBJID, PARAMCD, AVISITN", "COM.ADIS.HAI",
     "Source: sdtm is, adsl. HAI titers, three strains, Day 1 and Day 22."),
    ("ADFACE", "Reactogenicity Findings Analysis Dataset", "BASIC DATA STRUCTURE",
     "One record per subject per reaction per time point",
     "STUDYID, USUBJID, PARAMCD, ATPTN", "COM.ADFACE.GRADE",
     "Source: sdtm fa, vs, adsl. Daily diary grades, 30 minutes and Day 1-7."),
    ("ADCE", "Clinical Events Analysis Dataset", "OCCURRENCE DATA STRUCTURE",
     "One record per subject per solicited reaction per time window",
     "STUDYID, USUBJID, CEDECOD, ATPTN", "COM.ADCE.WINDOW",
     "Source: sdtm ce, adsl. Maximum grade within 30 minutes and within 7 days."),
    ("ADAE", "Adverse Events Analysis Dataset", "OCCURRENCE DATA STRUCTURE", "One record per adverse event per subject",
     "STUDYID, USUBJID, AESEQ", "COM.ADAE.WINDOW", "Source: sdtm ae, adsl. Unsolicited AEs and SAEs."),
]

# ---- ADSL variable catalog: name -> (label, type, length, codelist, role) ----
SL = OrderedDict([
    ("STUDYID", ("Study Identifier", "text", 12, "", "Identifier")),
    ("USUBJID", ("Unique Subject Identifier", "text", 20, "", "Identifier")),
    ("SUBJID", ("Subject Identifier for the Study", "text", 8, "", "Identifier")),
    ("SITEID", ("Study Site Identifier", "text", 8, "", "Record Qualifier")),
    ("AGE", ("Age", "integer", 8, "", "Record Qualifier")),
    ("AGEU", ("Age Units", "text", 5, "AGEU", "Variable Qualifier")),
    ("AGEGR1", ("Pooled Age Group 1", "text", 5, "AGEGR1", "Record Qualifier")),
    ("AGEGR1N", ("Pooled Age Group 1 (N)", "integer", 8, "AGEGR1N", "Variable Qualifier")),
    ("SEX", ("Sex", "text", 1, "SEX", "Record Qualifier")),
    ("RACE", ("Race", "text", 40, "RACE", "Record Qualifier")),
    ("ETHNIC", ("Ethnicity", "text", 30, "ETHNIC", "Record Qualifier")),
    ("COUNTRY", ("Country", "text", 3, "", "Record Qualifier")),
    ("TRT01P", ("Planned Treatment for Period 01", "text", 40, "ARM", "Record Qualifier")),
    ("TRT01PN", ("Planned Treatment for Period 01 (N)", "integer", 8, "TRTN", "Variable Qualifier")),
    ("TRT01A", ("Actual Treatment for Period 01", "text", 40, "ARM", "Record Qualifier")),
    ("TRT01AN", ("Actual Treatment for Period 01 (N)", "integer", 8, "TRTN", "Variable Qualifier")),
    ("ARM", ("Description of Planned Arm", "text", 40, "ARM", "Record Qualifier")),
    ("ACTARM", ("Description of Actual Arm", "text", 40, "ARM", "Record Qualifier")),
    ("ARMNRS", ("Reason Arm and/or Actual Arm is Null", "text", 40, "ARMNRS", "Record Qualifier")),
    ("RANDDT", ("Date of Randomization", "date", 8, "", "Timing")),
    ("TRTSDT", ("Date of First Exposure to Treatment", "date", 8, "", "Timing")),
    ("TRTSDTM", ("Datetime of First Exposure to Treatment", "datetime", 8, "", "Timing")),
    ("TRTEDT", ("Date of Last Exposure to Treatment", "date", 8, "", "Timing")),
    ("EOSDT", ("End of Study Date", "date", 8, "", "Timing")),
    ("EOSSTT", ("End of Study Status", "text", 12, "EOSSTT", "Record Qualifier")),
    ("DCSREAS", ("Reason for Discontinuation from Study", "text", 40, "DCSREAS", "Record Qualifier")),
    ("RANDFL", ("Randomized Population Flag", "text", 1, "NY", "Record Qualifier")),
    ("SAFFL", ("Safety Population Flag", "text", 1, "NY", "Record Qualifier")),
    ("IMMFL", ("Immunogenicity PP Population Flag", "text", 1, "NY", "Record Qualifier")),
])
SL_SRC = {  # origin, method
    "STUDYID": ("Predecessor", "", "DM.STUDYID"), "USUBJID": ("Predecessor", "", "DM.USUBJID"),
    "SUBJID": ("Predecessor", "", "DM.SUBJID"), "SITEID": ("Predecessor", "", "DM.SITEID"),
    "AGE": ("Predecessor", "", "DM.AGE"), "AGEU": ("Predecessor", "", "DM.AGEU"),
    "AGEGR1": ("Derived", "MT.ADSL.AGEGR", ""), "AGEGR1N": ("Derived", "MT.ADSL.AGEGR", ""),
    "SEX": ("Predecessor", "", "DM.SEX"), "RACE": ("Predecessor", "", "DM.RACE"),
    "ETHNIC": ("Predecessor", "", "DM.ETHNIC"), "COUNTRY": ("Predecessor", "", "DM.COUNTRY"),
    "ARM": ("Predecessor", "", "DM.ARM"), "ACTARM": ("Predecessor", "", "DM.ACTARM"),
    "ARMNRS": ("Predecessor", "", "DM.ARMNRS"),
    "TRT01P": ("Derived", "MT.ADSL.TRT", ""), "TRT01PN": ("Derived", "MT.ADSL.TRT", ""),
    "TRT01A": ("Derived", "MT.ADSL.TRT", ""), "TRT01AN": ("Derived", "MT.ADSL.TRT", ""),
    "RANDDT": ("Derived", "MT.ADSL.RANDDT", ""), "TRTSDT": ("Derived", "MT.ADSL.TRTDT", ""),
    "TRTSDTM": ("Derived", "MT.ADSL.TRTDT", ""), "TRTEDT": ("Derived", "MT.ADSL.TRTDT", ""),
    "EOSDT": ("Derived", "MT.ADSL.EOS", ""), "EOSSTT": ("Derived", "MT.ADSL.EOS", ""),
    "DCSREAS": ("Derived", "MT.ADSL.EOS", ""), "RANDFL": ("Derived", "MT.ADSL.RANDFL", ""),
    "SAFFL": ("Derived", "MT.ADSL.SAFFL", ""), "IMMFL": ("Derived", "MT.ADSL.IMMFL", ""),
}
SL_COM = {"TRTEDT": "COM.ADSL.TRTEDT", "IMMFL": "COM.ADSL.IMMFL"}

# variables carried into the other datasets from ADSL
CARRY = ["STUDYID", "USUBJID", "TRT01P", "TRT01PN", "TRT01A", "TRT01AN", "TRTSDT", "AGEGR1", "AGEGR1N", "SAFFL"]
CARRY_IMM = CARRY + ["IMMFL"]

A_, P_, D_ = "Assigned", "Predecessor", "Derived"

V = OrderedDict()
V["ADSL"] = [(k, *SL[k][:4], SL_SRC[k][0], "Sponsor", SL_SRC[k][1], SL[k][4], SL_COM.get(k, ""),
              SL_SRC[k][2]) for k in SL]


def carry(names):
    return [(k, *SL[k][:4], P_, "Sponsor", "", SL[k][4], "", f"ADSL.{k}") for k in names]


def row(var, label, typ, ln, cl="", origin=D_, meth="", role="", com="", note=""):
    return (var, label, typ, ln, cl, origin, "Sponsor", meth, role, com, note)


V["ADIS"] = carry(CARRY_IMM) + [
    row("PARAMCD", "Parameter Code", "text", 8, "PARAMIS", D_, "MT.ADIS.PARAM", "Topic"),
    row("PARAM", "Parameter", "text", 40, "", D_, "MT.ADIS.PARAM", "Synonym Qualifier"),
    row("PARAMN", "Parameter (N)", "integer", 8, "PARAMISN", D_, "MT.ADIS.PARAM", "Variable Qualifier"),
    row("PARCAT1", "Parameter Category 1", "text", 8, "", A_, "", "Grouping Qualifier", "", "HAI"),
    row("AVAL", "Analysis Value", "float", 8, "", D_, "MT.ADIS.AVAL", "Result Qualifier", "COM.ADIS.HAI"),
    row("AVALC", "Analysis Value (C)", "text", 8, "", P_, "", "Result Qualifier", "COM.ADIS.HAI", "IS.ISSTRESC"),
    row("LLOQ", "Lower Limit of Quantification", "integer", 8, "", P_, "", "Variable Qualifier", "", "IS.ISLLOQ"),
    row("ADT", "Analysis Date", "date", 8, "", D_, "MT.ADIS.DATE", "Timing"),
    row("ADY", "Analysis Relative Day", "integer", 8, "", D_, "MT.ADY", "Timing"),
    row("AVISIT", "Analysis Visit", "text", 8, "AVISIT", D_, "MT.ADIS.DATE", "Timing"),
    row("AVISITN", "Analysis Visit (N)", "integer", 8, "AVISITN", D_, "MT.ADIS.DATE", "Timing"),
    row("ABLFL", "Baseline Record Flag", "text", 1, "ABLFL", D_, "MT.ADIS.BASE", "Record Qualifier"),
    row("BASE", "Baseline Value", "float", 8, "", D_, "MT.ADIS.BASE", "Variable Qualifier", "COM.ADIS.HAI"),
    row("BASEC", "Baseline Value (C)", "text", 8, "", D_, "MT.ADIS.BASE", "Variable Qualifier"),
    row("BASECAT1", "Baseline Category 1", "text", 8, "BASECAT1", D_, "MT.ADIS.BASE", "Variable Qualifier"),
    row("R2BASE", "Ratio to Baseline", "float", 8, "", D_, "MT.ADIS.R2BASE", "Result Qualifier"),
    row("CRIT1", "Analysis Criterion 1", "text", 20, "CRIT1", A_, "", "Record Qualifier", "", "Titer >= 40"),
    row("CRIT1FL", "Criterion 1 Evaluation Result Flag", "text", 1, "NY", D_, "MT.ADIS.CRIT1", "Record Qualifier"),
    row("CRIT2", "Analysis Criterion 2", "text", 20, "CRIT2", A_, "", "Record Qualifier", "", "Seroconversion"),
    row("CRIT2FL", "Criterion 2 Evaluation Result Flag", "text", 1, "NY", D_, "MT.ADIS.CRIT2", "Record Qualifier"),
    row("ANL01FL", "Analysis Flag 01", "text", 1, "ANL01FL", D_, "MT.ADIS.ANL01FL", "Record Qualifier"),
    row("ASEQ", "Analysis Sequence Number", "integer", 8, "", D_, "MT.ASEQ", "Identifier"),
    row("SRCDOM", "Source Data", "text", 8, "", D_, "MT.SRC", "Record Qualifier", "", "IS"),
    row("SRCVAR", "Source Variable", "text", 8, "", D_, "MT.SRC", "Record Qualifier", "", "ISSTRESC"),
    row("SRCSEQ", "Source Sequence Number", "integer", 8, "", D_, "MT.SRC", "Record Qualifier"),
]

V["ADFACE"] = carry(CARRY) + [
    row("PARAMCD", "Parameter Code", "text", 8, "PARAMFA", D_, "MT.ADFACE.PARAM", "Topic"),
    row("PARAM", "Parameter", "text", 40, "", D_, "MT.ADFACE.PARAM", "Synonym Qualifier"),
    row("PARAMN", "Parameter (N)", "integer", 8, "PARAMFAN", D_, "MT.ADFACE.PARAM", "Variable Qualifier"),
    row("PARCAT1", "Parameter Category 1", "text", 8, "PARCAT1", D_, "MT.ADFACE.PARAM", "Grouping Qualifier"),
    row("AVAL", "Analysis Value", "integer", 8, "", D_, "MT.ADFACE.FA", "Result Qualifier", "COM.ADFACE.GRADE"),
    row("MEASVAL", "Measured Value", "float", 8, "", D_, "MT.ADFACE.FA", "Result Qualifier", "COM.ADFACE.MEAS"),
    row("MEASU", "Measured Value Unit", "text", 2, "", D_, "MT.ADFACE.FA", "Variable Qualifier", "COM.ADFACE.MEAS"),
    row("ADT", "Analysis Date", "date", 8, "", D_, "MT.ADFACE.TPT", "Timing"),
    row("ADY", "Analysis Relative Day", "integer", 8, "", D_, "MT.ADY", "Timing"),
    row("ATPT", "Analysis Timepoint", "text", 24, "ATPT", D_, "MT.ADFACE.TPT", "Timing"),
    row("ATPTN", "Analysis Timepoint (N)", "integer", 8, "ATPTN", D_, "MT.ADFACE.TPT", "Timing"),
    row("ASEQ", "Analysis Sequence Number", "integer", 8, "", D_, "MT.ASEQ", "Identifier"),
    row("SRCDOM", "Source Data", "text", 8, "", D_, "MT.ADFACE.SRC", "Record Qualifier"),
    row("SRCVAR", "Source Variable", "text", 8, "", D_, "MT.ADFACE.SRC", "Record Qualifier"),
    row("SRCSEQ", "Source Sequence Number", "integer", 8, "", D_, "MT.ADFACE.SRC", "Record Qualifier"),
]

V["ADCE"] = carry(CARRY) + [
    row("CESEQ", "Sequence Number", "integer", 8, "", P_, "", "Identifier", "", "CE.CESEQ"),
    row("CETERM", "Reported Term for the Clinical Event", "text", 20, "", P_, "", "Topic", "", "CE.CETERM"),
    row("CEDECOD", "Dictionary-Derived Term", "text", 20, "CEDECOD", P_, "", "Synonym Qualifier", "", "CE.CEDECOD"),
    row("CESCAT", "Subcategory for Clinical Event", "text", 8, "CESCAT", P_, "", "Grouping Qualifier", "", "CE.CESCAT"),
    row("CEOCCUR", "Clinical Event Occurrence", "text", 1, "NY", P_, "", "Record Qualifier", "", "CE.CEOCCUR"),
    row("ATOXGR", "Analysis Toxicity Grade", "text", 1, "ATOXGR", D_, "MT.ADCE.GRADE", "Record Qualifier",
        "COM.ADFACE.GRADE"),
    row("ATOXGRN", "Analysis Toxicity Grade (N)", "integer", 8, "ATOXGRN", D_, "MT.ADCE.GRADE", "Variable Qualifier"),
    row("ATPT", "Analysis Timepoint", "text", 24, "ATPT", D_, "MT.ADCE.ATPT", "Timing"),
    row("ATPTN", "Analysis Timepoint (N)", "integer", 8, "ATPTN", D_, "MT.ADCE.ATPT", "Timing"),
    row("ASTDT", "Analysis Start Date", "date", 8, "", D_, "MT.ADCE.DATES", "Timing"),
    row("AENDT", "Analysis End Date", "date", 8, "", D_, "MT.ADCE.DATES", "Timing"),
    row("ASTDY", "Analysis Start Relative Day", "integer", 8, "", D_, "MT.ADY", "Timing"),
    row("AENDY", "Analysis End Relative Day", "integer", 8, "", D_, "MT.ADY", "Timing"),
    row("AOCCFL", "1st Occurrence within Subject Flag", "text", 1, "NY", D_, "MT.ADCE.OCC", "Record Qualifier"),
    row("AOCCSFL", "1st Occurrence of SOC Flag", "text", 1, "NY", D_, "MT.ADCE.OCC", "Record Qualifier"),
    row("AOCCPFL", "1st Occurrence of Preferred Term Flag", "text", 1, "NY", D_, "MT.ADCE.OCC", "Record Qualifier"),
]

V["ADAE"] = carry(CARRY) + [
    row("AESEQ", "Sequence Number", "integer", 8, "", P_, "", "Identifier", "", "AE.AESEQ"),
    row("AETERM", "Reported Term for the Adverse Event", "text", 40, "", P_, "", "Topic", "", "AE.AETERM"),
    row("AEDECOD", "Dictionary-Derived Term", "text", 40, "", P_, "", "Synonym Qualifier", "COM.ADAE.CODING",
        "AE.AEDECOD"),
    row("AEBODSYS", "Body System or Organ Class", "text", 60, "", P_, "", "Record Qualifier", "COM.ADAE.CODING",
        "AE.AEBODSYS"),
    row("AECAT", "Category for Adverse Event", "text", 12, "", P_, "", "Grouping Qualifier", "", "AE.AECAT"),
    row("AESEV", "Severity/Intensity", "text", 8, "AESEV", P_, "", "Record Qualifier", "", "AE.AESEV"),
    row("AESEVN", "Severity/Intensity (N)", "integer", 8, "AESEVN", D_, "MT.ADAE.SEV", "Variable Qualifier"),
    row("AESER", "Serious Event", "text", 1, "NY", P_, "", "Record Qualifier", "", "AE.AESER"),
    row("AEREL", "Causality", "text", 20, "", P_, "", "Record Qualifier", "", "AE.AEREL"),
    row("AEACN", "Action Taken with Study Treatment", "text", 20, "", P_, "", "Record Qualifier", "", "AE.AEACN"),
    row("AEOUT", "Outcome of Adverse Event", "text", 30, "", P_, "", "Record Qualifier", "", "AE.AEOUT"),
    row("AESTDTC", "Start Date/Time of Adverse Event", "text", 10, "", P_, "", "Timing", "", "AE.AESTDTC"),
    row("AEENDTC", "End Date/Time of Adverse Event", "text", 10, "", P_, "", "Timing", "", "AE.AEENDTC"),
    row("ASTDT", "Analysis Start Date", "date", 8, "", D_, "MT.ADAE.DATES", "Timing"),
    row("ASTDTF", "Analysis Start Date Imputation Flag", "text", 1, "ASTDTF", D_, "MT.ADAE.DATES", "Timing"),
    row("AENDT", "Analysis End Date", "date", 8, "", D_, "MT.ADAE.DATES", "Timing"),
    row("ASTDY", "Analysis Start Relative Day", "integer", 8, "", D_, "MT.ADY", "Timing"),
    row("AENDY", "Analysis End Relative Day", "integer", 8, "", D_, "MT.ADY", "Timing"),
    row("TRTEMFL", "Treatment Emergent Analysis Flag", "text", 1, "NY", D_, "MT.ADAE.TRTEM", "Record Qualifier"),
    row("ANL01FL", "Analysis Flag 01", "text", 1, "ANL01FL", D_, "MT.ADAE.WIN", "Record Qualifier",
        "COM.ADAE.WINDOW"),
    row("ANL02FL", "Analysis Flag 02", "text", 1, "ANL01FL", D_, "MT.ADAE.WIN", "Record Qualifier",
        "COM.ADAE.WINDOW"),
    row("AOCCFL", "1st Occurrence within Subject Flag", "text", 1, "NY", D_, "MT.ADAE.OCC", "Record Qualifier"),
    row("AOCCSFL", "1st Occurrence of SOC Flag", "text", 1, "NY", D_, "MT.ADAE.OCC", "Record Qualifier"),
    row("AOCCPFL", "1st Occurrence of Preferred Term Flag", "text", 1, "NY", D_, "MT.ADAE.OCC", "Record Qualifier"),
]

KEY_MAND = {"STUDYID", "USUBJID", "PARAMCD", "PARAM", "AESEQ", "CESEQ"}
var_rows = []
for ds, rows in V.items():
    for i, (var, label, typ, ln, cl, origin, src, meth, role, com, note) in enumerate(rows, 1):
        assigned = note if origin == A_ else None
        dnote = None if origin == A_ else (note or None)
        var_rows.append((i, ds, var, label, typ, ln, "Yes" if var in KEY_MAND else "No", assigned, cl or None,
                         origin, src, None, meth or None, role, com or None, dnote))

# ---- value level ----
vl_rows = []
n = 0
for pcds, unit in [("REDNESS, SWELLING, INDURAT", "mm"), ("FEVER", "C")]:
    n += 1
    vl_rows.append((n, "ADFACE", "MEASU", f"PARAMCD IN ({pcds})", "text", 2, "", D_, "Sponsor", None,
                    "MT.ADFACE.FA", unit))

# ---- codelists ----
sdtm_wb = openpyxl.load_workbook(SDTM_SPEC)
sd_cl = [r for r in sdtm_wb["Codelists"].iter_rows(min_row=2, values_only=True)]
cl_rows = []


def from_sdtm(src_id, new_id=None, name=None, terms=None):
    k = 0
    for r in sd_cl:
        if r[0] == src_id and (terms is None or r[6] in terms):
            k += 1
            cl_rows.append((new_id or src_id, name or r[1], r[2], r[3], r[4], k, r[6], r[7], r[8]))


def custom(cid, name, typ, terms):
    for k, t in enumerate(terms, 1):
        term, dec = t if isinstance(t, tuple) else (t, None)
        cl_rows.append((cid, name, None, typ, "Sponsor defined", k, str(term), None, dec))


from_sdtm("NY")
from_sdtm("AGEU")
from_sdtm("SEX")
from_sdtm("RACE")
from_sdtm("ETHNIC")
from_sdtm("AESEV")
from_sdtm("ARMNULRS", "ARMNRS", "Reason Arm and/or Actual Arm is Null")
from_sdtm("DSDECOD", "DCSREAS", "Reason for Discontinuation from Study",
          ["LOST TO FOLLOW-UP", "SCREEN FAILURE", "WITHDRAWAL BY SUBJECT"])
custom("ARM", "Description of Arm", "text", ["VAXF-101 0.5 mL", "Placebo"])
custom("TRTN", "Treatment (N)", "integer", [("1", "VAXF-101 0.5 mL"), ("2", "Placebo")])
custom("AGEGR1", "Pooled Age Group 1", "text", ["18-45", "46-60"])
custom("AGEGR1N", "Pooled Age Group 1 (N)", "integer", [("1", "18-45"), ("2", "46-60")])
custom("EOSSTT", "End of Study Status", "text", ["COMPLETED", "DISCONTINUED"])
custom("PARAMIS", "Immunogenicity Parameter", "text",
       [("HAIA", "HAI Titer, Strain A"), ("HAIB", "HAI Titer, Strain B"), ("HAIC", "HAI Titer, Strain C")])
custom("PARAMISN", "Immunogenicity Parameter (N)", "integer",
       [("1", "HAI Titer, Strain A"), ("2", "HAI Titer, Strain B"), ("3", "HAI Titer, Strain C")])
custom("AVISIT", "Analysis Visit", "text", ["DAY 1", "DAY 22"])
custom("AVISITN", "Analysis Visit (N)", "integer", [("2", "DAY 1"), ("4", "DAY 22")])
custom("ABLFL", "Baseline Record Flag", "text", [("Y", "Yes")])
custom("ANL01FL", "Analysis Flag", "text", [("Y", "Yes")])
custom("BASECAT1", "Baseline Category 1", "text", [("<10", "Seronegative"), (">=10", "Seropositive")])
custom("CRIT1", "Analysis Criterion 1", "text", ["Titer >= 40"])
custom("CRIT2", "Analysis Criterion 2", "text", ["Seroconversion"])

REACT = [("PAIN", "Pain", "LOCAL"), ("TENDERN", "Tenderness", "LOCAL"), ("REDNESS", "Redness", "LOCAL"),
         ("SWELLING", "Swelling", "LOCAL"), ("INDURAT", "Induration", "LOCAL"), ("FEVER", "Fever", "SYSTEMIC"),
         ("CHILLS", "Chills", "SYSTEMIC"), ("MALAISE", "Malaise", "SYSTEMIC"), ("MYALGIA", "Myalgia", "SYSTEMIC"),
         ("HEADACHE", "Headache", "SYSTEMIC"), ("ARTHRALG", "Arthralgia", "SYSTEMIC"),
         ("NAUSEA", "Nausea", "SYSTEMIC"), ("VOMITING", "Vomiting", "SYSTEMIC")]
custom("PARAMFA", "Reactogenicity Parameter", "text", [(c, f"{p} Grade") for c, p, _ in REACT])
custom("PARAMFAN", "Reactogenicity Parameter (N)", "integer", [(str(i), f"{p} Grade") for i, (_, p, _) in
                                                                enumerate(REACT, 1)])
custom("PARCAT1", "Parameter Category 1", "text", ["LOCAL", "SYSTEMIC"])
custom("CEDECOD", "Reaction Term", "text", [
    "PAIN", "TENDERNESS", "REDNESS", "SWELLING", "INDURATION", "FEVER", "CHILLS", "MALAISE", "MYALGIA", "HEADACHE",
    "ARTHRALGIA", "NAUSEA", "VOMITING"])
custom("CESCAT", "Reaction Subcategory", "text", ["LOCAL", "SYSTEMIC"])
custom("ATPT", "Analysis Timepoint", "text",
       ["30 MINUTES POST-DOSE"] + [f"DAY {d}" for d in range(1, 8)] + ["WITHIN 7 DAYS POST-DOSE"])
custom("ATPTN", "Analysis Timepoint (N)", "integer",
       [("0", "30 MINUTES POST-DOSE")] + [(str(d), f"DAY {d}") for d in range(1, 8)] + [("8", "WITHIN 7 DAYS POST-DOSE")])
custom("ATOXGR", "Analysis Toxicity Grade", "text",
       [("0", "None"), ("1", "Grade 1"), ("2", "Grade 2"), ("3", "Grade 3"), ("4", "Grade 4")])
custom("ATOXGRN", "Analysis Toxicity Grade (N)", "integer",
       [("0", "None"), ("1", "Grade 1"), ("2", "Grade 2"), ("3", "Grade 3"), ("4", "Grade 4")])
custom("ASTDTF", "Date Imputation Flag", "text", [("D", "Day imputed")])
custom("AESEVN", "Severity/Intensity (N)", "integer", [("1", "MILD"), ("2", "MODERATE"), ("3", "SEVERE")])

# ---- methods ----
meth = [
    ("MT.ADY", "Algorithm", "Relative day: date - TRTSDT + 1 if date >= TRTSDT, else date - TRTSDT. Null if either is null."),
    ("MT.ASEQ", "Algorithm", "Running number within USUBJID, ordered PARAMN, then time point or visit."),
    ("MT.SRC", "Algorithm", "SRCDOM, SRCVAR, SRCSEQ point to the IS record the value comes from."),
    ("MT.ADSL.AGEGR", "Algorithm", "AGE <= 45: '18-45' (AGEGR1N 1); AGE >= 46: '46-60' (2). Matches the randomization stratum."),
    ("MT.ADSL.TRT", "Algorithm",
     "TRT01P = DM.ARM, TRT01A = DM.ACTARM. TRT01PN, TRT01AN: VAXF-101 1, Placebo 2. Null for screen failures; "
     "TRT01A also null for the 2 randomized subjects not dosed."),
    ("MT.ADSL.RANDDT", "Algorithm", "Date part of DS.DSSTDTC where DS.DSDECOD = RANDOMIZED."),
    ("MT.ADSL.TRTDT", "Algorithm",
     "TRTSDTM = DM.RFXSTDTC as datetime, TRTSDT = its date part. TRTEDT = date part of DM.RFXENDTC. "
     "Single dose, so TRTEDT = TRTSDT. Null if not dosed."),
    ("MT.ADSL.EOS", "Algorithm",
     "From the DS record with DSCAT = DISPOSITION EVENT (one per subject). EOSSTT = COMPLETED if DSDECOD = COMPLETED, "
     "else DISCONTINUED. EOSDT = date part of DSSTDTC. DCSREAS = DSDECOD when EOSSTT = DISCONTINUED. "
     "Screen failures: DISCONTINUED, SCREEN FAILURE."),
    ("MT.ADSL.RANDFL", "Algorithm", "'Y' if a DS record with DSDECOD = RANDOMIZED exists, else 'N'."),
    ("MT.ADSL.SAFFL", "Algorithm", "'Y' if RANDFL = 'Y' and the subject has an EX record (TRTSDT not null), else 'N'."),
    ("MT.ADSL.IMMFL", "Algorithm",
     "'Y' if SAFFL = 'Y' and all three strains have a valid HAI result (ISSTAT null) at Day 1 and Day 22, and the "
     "Day 22 sample is within window (study day 19 to 25, TRTSDT = day 1). Otherwise 'N'."),
    ("MT.ADIS.PARAM", "Algorithm", "ISBDAGNT STRAIN A/B/C -> PARAMCD HAIA/HAIB/HAIC, PARAMN 1/2/3. PARAM per codelist PARAMIS."),
    ("MT.ADIS.AVAL", "Algorithm",
     "AVALC = IS.ISSTRESC. AVAL = IS.ISSTRESN; if ISSTRESC = '<10', AVAL = LLOQ / 2 = 5. Null if ISSTAT = NOT DONE."),
    ("MT.ADIS.DATE", "Algorithm", "ADT = date part of IS.ISDTC. AVISIT = IS.VISIT, AVISITN = IS.VISITNUM."),
    ("MT.ADIS.BASE", "Algorithm",
     "ABLFL = 'Y' on the Day 1 record where IS.ISLOBXFL = 'Y' and AVAL is not null. BASE, BASEC = AVAL, AVALC of that "
     "record, carried to every record of the same subject and PARAMCD. BASECAT1 = '<10' if BASEC = '<10', else '>=10'."),
    ("MT.ADIS.R2BASE", "Algorithm", "AVAL / BASE on Day 22 records. Null on Day 1 records or if AVAL or BASE is null."),
    ("MT.ADIS.CRIT1", "Algorithm", "CRIT1FL = 'Y' if AVAL >= 40, 'N' if AVAL < 40, null if AVAL is null. Day 1 and Day 22."),
    ("MT.ADIS.CRIT2", "Algorithm",
     "Day 22 only. Seroconversion: BASECAT1 = '<10' and AVAL >= 40, or BASECAT1 = '>=10' and R2BASE >= 4. "
     "CRIT2FL = 'Y' if met, 'N' if not, null if AVAL or BASE is null. CRIT2 and CRIT2FL are null on Day 1."),
    ("MT.ADIS.ANL01FL", "Algorithm", "'Y' if IMMFL = 'Y' and AVAL is not null."),
    ("MT.ADFACE.PARAM", "Algorithm",
     "One parameter per reaction: PARAMCD = reaction (codelist PARAMFA), PARCAT1 = LOCAL or SYSTEMIC. "
     "Local: pain, tenderness, redness, swelling, induration. Systemic: fever and the other 7."),
    ("MT.ADFACE.FA", "Algorithm",
     "Reactions other than fever, per FA.FAOBJ and FATPT. Redness, swelling, induration: MEASVAL = FA DIAMETER "
     "FASTRESN (mm), AVAL = 0 if < 25, 1 if 25-50, 2 if 51-100, 3 if > 100. Pain, tenderness, chills, malaise, myalgia, "
     "headache, arthralgia, nausea, vomiting: AVAL = 0 if FA OCCUR = N, else 1/2/3 for FA SEV = MILD/MODERATE/SEVERE; "
     "MEASVAL null. Fever (VS TEMP, VSCAT = REACTOGENICITY, VSSTRESN in C): MEASVAL = VSSTRESN, AVAL = 0 if < 38.0, "
     "1 if 38.0-38.4, 2 if 38.5-38.9, 3 if 39.0-40.0, 4 if > 40.0. MEASU = mm or C."),
    ("MT.ADFACE.TPT", "Algorithm",
     "FA.FATPT, VS.VSTPT -> ATPT (30 MINUTES POST-DOSE, DAY 1-DAY 7), ATPTN 0-7. ADT = date part of FADTC or VSDTC."),
    ("MT.ADFACE.SRC", "Algorithm",
     "SRCDOM FA or VS. SRCVAR and SRCSEQ point to the record that gives AVAL: FA DIAMETER (FASTRESN) for diameter "
     "reactions; FA SEV (FASTRESC) when FA OCCUR = Y, else the OCCUR record, for the others; VS (VSSTRESN) for fever."),
    ("MT.ADCE.GRADE", "Algorithm",
     "CEOCCUR = N: ATOXGR '0'. Else CETOXGR if not null (diameter and temperature reactions), else CESEV "
     "MILD/MODERATE/SEVERE -> '1'/'2'/'3'. ATOXGRN = numeric ATOXGR."),
    ("MT.ADCE.ATPT", "Algorithm",
     "CETPT '30 MINUTES POST-DOSE' -> ATPT the same, ATPTN 0. CETPT 'DAY 7' -> ATPT 'WITHIN 7 DAYS POST-DOSE', ATPTN 8."),
    ("MT.ADCE.DATES", "Algorithm", "ASTDT, AENDT = date part of CESTDTC, CEENDTC. Null when no reaction."),
    ("MT.ADCE.OCC", "Algorithm",
     "Among records with CEOCCUR = Y, within USUBJID and ATPTN, ordered by ATOXGRN descending then CESEQ: AOCCFL = 'Y' "
     "on the first record; AOCCSFL on the first within CESCAT; AOCCPFL on every such record (one per reaction). "
     "Null otherwise. The flagged record carries the maximum grade."),
    ("MT.ADAE.DATES", "Algorithm",
     "AENDT = date part of AEENDTC. ASTDT = date part of AESTDTC. If AESTDTC is year-month only: ASTDT = TRTSDT when it "
     "falls in the same month, else the first day of that month; ASTDTF = 'D'. Otherwise ASTDTF is null."),
    ("MT.ADAE.SEV", "Algorithm", "AESEV MILD 1, MODERATE 2, SEVERE 3."),
    ("MT.ADAE.TRTEM", "Algorithm", "'Y' if ASTDT >= TRTSDT, else null. Null if ASTDT or TRTSDT is null."),
    ("MT.ADAE.WIN", "Algorithm",
     "ANL01FL = 'Y' if AECAT = UNSOLICITED, TRTEMFL = 'Y' and ASTDY <= 22. "
     "ANL02FL = 'Y' if AESER = 'Y' and TRTEMFL = 'Y' and ASTDY <= 91."),
    ("MT.ADAE.OCC", "Algorithm",
     "Among records with TRTEMFL = 'Y', within USUBJID: AOCCFL = 'Y' on the first record by ASTDT, AESEQ; "
     "AOCCSFL on the first within AEBODSYS; AOCCPFL on the first within AEBODSYS and AEDECOD."),
]

com = [
    ("COM.ADSL.POP", "ADSL has all 267 screened subjects, including 27 screen failures (RANDFL = N). Populations: "
                     "safety = randomized and dosed (238); immunogenicity per protocol = IMMFL."),
    ("COM.ADSL.TRTEDT", "Single 0.5 mL dose on Day 1, so TRTEDT = TRTSDT."),
    ("COM.ADSL.IMMFL", "Window and valid-result rules follow docs/study-design.md (Day 22 +/- 3 days)."),
    ("COM.ADIS.HAI", "HAI titers are reciprocal dilutions. Results below LLOQ (10) are reported '<10' and enter "
                     "GMT, GMFR and seroconversion as LLOQ / 2 = 5. AVALC keeps the reported text. GMT is computed "
                     "in the TLF program as exp(mean(ln AVAL))."),
    ("COM.ADFACE.GRADE", "FDA 2007 toxicity grading scale for preventive vaccine trials. Grade 4 is possible only for "
                         "fever (> 40.0 C); the other reactions have no ER-visit or necrosis data, so they stop at 3. "
                         "Missing diary days have no record; no imputation."),
    ("COM.ADFACE.MEAS", "MEASVAL is diameter (mm) for redness, swelling, induration and temperature (C) for fever. "
                        "Null for severity-scale reactions."),
    ("COM.ADCE.WINDOW", "SDTM CE holds one summary record per reaction per window. ADCE restates the grade from CE. "
                        "Cross-check: ATOXGRN equals the maximum ADFACE AVAL for the same reaction and window "
                        "(30 minutes; Day 1-7)."),
    ("COM.ADAE.WINDOW", "Unsolicited AEs are counted through Day 22 (ANL01FL), SAEs through Day 91 (ANL02FL)."),
    ("COM.ADAE.CODING", "No MedDRA. AEDECOD and AEBODSYS are carried from SDTM, which uses a local coding table."),
]

issues = [
    (1, "ADSL", "Day 22 window (19-25) excludes 5 subjects whose sample was drawn outside it. They are in the "
                "safety population but not in IMMFL.", "Closed"),
    (2, "ADIS", "BASE and seroconversion use AVAL = 5 for '<10'. Baseline '<10' (BASECAT1) selects the "
                "seroconversion rule, not BASE itself.", "Closed"),
    (3, "ADIS", "Subjects with a NOT DONE titer keep their record with null AVAL and are outside ANL01FL.", "Closed"),
    (4, "ADFACE", "Fever comes from VS (diary temperature), not FACE. Redness, swelling, induration diameters under "
                  "25 mm are grade 0.", "Closed"),
    (5, "ADFACE", "SDTM CE/FACE are derived in SDTM (CETOXGR); ADCE and ADFACE grade independently so the two can be "
                  "cross-checked.", "Closed"),
    (6, "ADAE", "No MedDRA version. SOC and PT are the local table values.", "Closed"),
    (7, "ADAE", "Two AEs (101-030, 101-096) start in the same month as the dose, year-month only. ASTDT is set to the "
                "dose date, so they count as treatment-emergent and fall inside the Day 22 window.", "Closed"),
    (8, "All", "Compare ADSL and ADIS against {admiralvaccine} in R after SAS = R QC. Skipped if time is short.",
     "Open"),
]

wb = openpyxl.Workbook()
wb.remove(wb.active)
HFONT, HFILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")


def sheet(title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    for c in ws[1]:
        c.font, c.fill = HFONT, HFILL
    for rw in ws.iter_rows(min_row=2):
        for c in rw:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"


sheet("Study", ["Attribute", "Value"], [
    ("StudyName", "VAXF101"),
    ("StudyDescription", "Synthetic randomized, double-blind, placebo-controlled study of an influenza vaccine, "
                         "ADaM datasets (ADaMIG 1.3)"),
    ("ProtocolName", "VAXF101"),
    ("Language", "en"),
    ("Scope", "ADSL, ADIS, ADFACE, ADCE, ADAE. Learning/portfolio use only; not related to any regulatory "
              "submission. Source: P3 SDTM datasets built from synthetic raw data (python/make_raw.py).")],
      [20, 100])
sheet("Standards", ["OID", "Name", "Type", "Publishing Set", "Version", "Status"], [
    ("STD.ADAMIG", "ADaMIG", "IG", "", "1.3", "Final"),
    ("STD.CT.SDTM", "CDISC/NCI", "CT", "SDTM", "2026-09-25", "Final"),
    ("STD.CT.SDTM.2025", "CDISC/NCI", "CT", "SDTM", "2025-09-26", "Final")], [20, 14, 8, 14, 12, 10])
sheet("Datasets", ["Dataset", "Description", "Class", "Structure", "Purpose", "Key Variables", "Standard", "Comment",
                   "Developer Notes"],
      [(d, desc, cls, st, "Analysis", keys, "ADaMIG 1.3", c, note) for d, desc, cls, st, keys, c, note in datasets],
      [9, 30, 24, 40, 10, 34, 12, 20, 60])
sheet("Variables", ["Order", "Dataset", "Variable", "Label", "Data Type", "Length", "Mandatory", "Assigned Value",
                    "Codelist", "Origin", "Source", "Pages", "Method", "Role", "Comment", "Developer Notes"],
      var_rows, [6, 8, 10, 36, 9, 7, 9, 13, 11, 11, 9, 7, 16, 17, 16, 20])
sheet("ValueLevel", ["Order", "Dataset", "Variable", "Where Clause", "Data Type", "Length", "Codelist", "Origin",
                     "Source", "Pages", "Method", "Developer Notes"], vl_rows,
      [6, 8, 9, 40, 10, 7, 10, 10, 9, 7, 16, 14])
sheet("Codelists", ["ID", "Name", "NCI Codelist Code", "Data Type", "Terminology", "Order", "Term", "NCI Term Code",
                    "Decoded Value"], cl_rows, [10, 34, 12, 9, 24, 6, 34, 12, 30])
sheet("Methods", ["ID", "Type", "Description"], meth, [20, 10, 120])
sheet("Comments", ["ID", "Description"], com, [20, 120])
sheet("Open Issues", ["#", "Dataset", "Issue", "Status"], issues, [4, 10, 110, 8])
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)

# ---- consistency checks ----
mt_ids = {m[0] for m in meth}
com_ids = {c[0] for c in com}
cl_ids = {c[0] for c in cl_rows}
for r in var_rows:
    assert not r[12] or r[12] in mt_ids, r
    assert not r[14] or r[14] in com_ids, r
    assert not r[8] or r[8] in cl_ids, r
for r in vl_rows:
    assert not r[10] or r[10] in mt_ids, r
for d, *_ in datasets:
    assert d in V, d
used = {r[12] for r in var_rows if r[12]} | {r[10] for r in vl_rows}
print("variables:", len(var_rows), "| value level:", len(vl_rows), "| codelist terms:", len(cl_rows))
print("methods unused:", sorted(mt_ids - used))
print("comments unused:", sorted(com_ids - {r[14] for r in var_rows if r[14]} - {d[5] for d in datasets}
                                 - {c for c in com_ids if c in {r[14] for r in var_rows}}))
