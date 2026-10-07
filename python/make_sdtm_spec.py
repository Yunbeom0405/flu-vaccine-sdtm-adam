"""Build the P3 SDTM spec workbook (same layout as P2)."""
import csv
from collections import OrderedDict
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "projects/P3-vaccine-immuno/specs/sdtm-spec.xlsx"
IG = ROOT / "reference/cdisc/sdtmig34-variables.csv"
CT_XLS = ROOT / "projects/P1-cdisc-pilot-e2e/SDTM Terminology.xls"
RACE_CSV = ROOT / "reference/p2-spec-build/race_cl.csv"
CT = "CDISC/NCI CT 2026-09-25"
CT_OLD = "CDISC/NCI CT 2025-09-26"

ig = {}
for x in csv.DictReader(open(IG)):
    ig.setdefault((x["domain"], x["name"]), x)

ct = pd.read_excel(CT_XLS, sheet_name=1)
ct.columns = ["code", "cl", "ext", "clname", "val", "syn", "defn", "pt"]
heads = ct[ct["cl"].isna()].set_index("code")

# timing variables missing from the parsed IG table (checked later by P21)
TIMING = {
    "VISITNUM": ("Visit Number", "Timing"), "VISIT": ("Visit Name", "Timing"),
    "VISITDY": ("Planned Study Day of Visit", "Timing"), "EPOCH": ("Epoch", "Timing"),
    "FATPT": ("Planned Time Point Name", "Timing"), "FATPTNUM": ("Planned Time Point Number", "Timing"),
    "FATPTREF": ("Time Point Reference", "Timing"), "FARFTDTC": ("Date/Time of Reference Time Point", "Timing"),
    "CETPT": ("Planned Time Point Name", "Timing"), "CETPTNUM": ("Planned Time Point Number", "Timing"),
    "CETPTREF": ("Time Point Reference", "Timing"), "CERFTDTC": ("Date/Time of Reference Time Point", "Timing"),
    "CEEVINTX": ("Evaluation Interval Text", "Timing"), "FAEVINTX": ("Evaluation Interval Text", "Timing"),
}

datasets = [
    ("DM", "Demographics", "SPECIAL PURPOSE", "One record per subject", "STUDYID, USUBJID", "",
     "All 267 screened subjects, 27 screen failures. Source: DM, IC, RAND, EX, DS."),
    ("DS", "Disposition", "EVENTS", "One record per disposition status or protocol milestone per subject",
     "STUDYID, USUBJID, DSDECOD, DSSTDTC", "", "Source: IC, RAND, DS."),
    ("EX", "Exposure", "INTERVENTIONS", "One record per constant dosing interval per subject",
     "STUDYID, USUBJID, EXTRT, EXSTDTC", "", "Source: EX, RAND. Only subjects who were dosed."),
    ("SV", "Subject Visits", "SPECIAL PURPOSE", "One record per actual visit per subject",
     "STUDYID, USUBJID, VISITNUM", "", "Source: VISIT."),
    ("IE", "Inclusion/Exclusion Criterion Not Met", "FINDINGS",
     "One record per inclusion/exclusion criterion not met per subject",
     "STUDYID, USUBJID, IETESTCD", "", "Source: IE. Screen failures only."),
    ("VS", "Vital Signs", "FINDINGS", "One record per vital sign measurement per time point per visit per subject",
     "STUDYID, USUBJID, VSTESTCD, VISITNUM, VSDTC, VSTPTNUM", "COM.TEMP",
     "Source: VITALS (screening, Day 1) and DIARY (daily temperature)."),
    ("CE", "Clinical Events", "EVENTS", "One record per solicited reaction per time point per subject",
     "STUDYID, USUBJID, CEDECOD, CETPTNUM", "COM.FLAT",
     "Solicited reactions (global record per reaction, flat model). Source: DIARY, OBS30."),
    ("FA", "Findings About Events or Interventions", "FINDINGS",
     "One record per finding per object per time point per visit per subject",
     "STUDYID, USUBJID, FAOBJ, FATESTCD, FATPTNUM", "COM.FLAT",
     "FACE. Daily diary values. Source: DIARY, OBS30."),
    ("AE", "Adverse Events", "EVENTS", "One record per adverse event per subject",
     "STUDYID, USUBJID, AEDECOD, AESTDTC", "", "Unsolicited AEs and SAEs. Source: AE."),
    ("IS", "Immunogenicity Specimen Assessments", "FINDINGS",
     "One record per test per specimen per visit per subject", "STUDYID, USUBJID, ISTESTCD, ISBDAGNT, VISITNUM", "",
     "HAI titers, three strains. Source: LB_HAI."),
    ("TA", "Trial Arms", "TRIAL DESIGN", "One record per planned element per arm", "STUDYID, ARMCD, TAETORD", "COM.TD", ""),
    ("TE", "Trial Elements", "TRIAL DESIGN", "One record per planned element", "STUDYID, ETCD", "COM.TD", ""),
    ("TV", "Trial Visits", "TRIAL DESIGN", "One record per planned visit per arm", "STUDYID, ARMCD, VISITNUM", "COM.TD", ""),
    ("TI", "Trial Inclusion/Exclusion Criteria", "TRIAL DESIGN", "One record per I/E criterion", "STUDYID, IETESTCD",
     "COM.TD", ""),
    ("TS", "Trial Summary", "TRIAL DESIGN", "One record per trial summary parameter value",
     "STUDYID, TSPARMCD, TSSEQ", "COM.TD", ""),
]

C, D, A = "Collected", "Derived", "Assigned"


def s(form, var):
    return f"{form}.{var}"


V = OrderedDict()
V["DM"] = [
    ("STUDYID", A, "Sponsor", "", "", "", ""),
    ("DOMAIN", A, "Sponsor", "", "", "", ""),
    ("USUBJID", D, "Sponsor", "", "MT.USUBJID", "", s("DM", "SUBJECT")),
    ("SUBJID", C, "Sponsor", "", "", "", s("DM", "SUBJECT")),
    ("RFSTDTC", D, "Sponsor", "", "MT.RFDOSE", "", "First dose date/time"),
    ("RFENDTC", D, "Sponsor", "", "MT.RFEND", "", "Last disposition date"),
    ("RFXSTDTC", D, "Sponsor", "", "MT.RFDOSE", "", "First dose date/time"),
    ("RFXENDTC", D, "Sponsor", "", "MT.RFDOSE", "", "Single dose, same as RFXSTDTC"),
    ("RFICDTC", C, "Investigator", "", "MT.ISODT", "", s("IC", "ICDAT")),
    ("RFPENDTC", D, "Sponsor", "", "MT.RFEND", "", "Same as RFENDTC"),
    ("SITEID", C, "Sponsor", "", "", "", s("DM", "SITE")),
    ("BRTHDTC", C, "Investigator", "", "MT.ISODT", "", s("DM", "BRTHDAT")),
    ("AGE", D, "Sponsor", "", "MT.AGE", "", "At RFICDTC"),
    ("AGEU", A, "Sponsor", "AGEU", "", "", ""),
    ("SEX", C, "Investigator", "SEX", "MT.UPPER", "", s("DM", "SEX")),
    ("RACE", C, "Investigator", "RACE", "MT.UPPER", "COM.RACE", s("DM", "RACE")),
    ("ETHNIC", C, "Investigator", "ETHNIC", "MT.UPPER", "COM.RACE", s("DM", "ETHNIC")),
    ("ARMCD", D, "Sponsor", "", "MT.ARM", "", s("RAND", "TRTCODE")),
    ("ARM", D, "Sponsor", "", "MT.ARM", "", s("RAND", "TRTCODE")),
    ("ACTARMCD", D, "Sponsor", "", "MT.ARM", "", "Null if not dosed"),
    ("ACTARM", D, "Sponsor", "", "MT.ARM", "", "Null if not dosed"),
    ("ARMNRS", D, "Sponsor", "ARMNULRS", "MT.ARM", "COM.ARMNRS", ""),
    ("COUNTRY", A, "Sponsor", "", "", "", "ISO 3166-1 alpha-3. Synthetic, one country"),
    ("DMDTC", C, "Investigator", "", "MT.ISODT", "", s("IC", "ICDAT")),
    ("DMDY", D, "Sponsor", "", "MT.DY", "", ""),
]
V["DS"] = [
    ("STUDYID", A, "Sponsor", "", "", "", ""),
    ("DOMAIN", A, "Sponsor", "", "", "", ""),
    ("USUBJID", D, "Sponsor", "", "MT.USUBJID", "", ""),
    ("DSSEQ", D, "Sponsor", "", "MT.SEQ", "", ""),
    ("DSREFID", C, "Sponsor", "", "", "", "RAND.RANDNUM, on the RANDOMIZED record"),
    ("DSTERM", C, "Investigator", "", "", "", "DS.DSTERM; milestone text assigned"),
    ("DSDECOD", D, "Sponsor", "DSDECOD", "MT.DSDECOD", "COM.DS", ""),
    ("DSCAT", A, "Sponsor", "DSCAT", "", "", ""),
    ("EPOCH", D, "Sponsor", "EPOCH", "MT.EPOCH", "", ""),
    ("DSSTDTC", C, "Investigator", "", "MT.ISODT", "", "IC.ICDAT, RAND.RANDDAT, DS.DSDAT"),
    ("DSSTDY", D, "Sponsor", "", "MT.DY", "", ""),
]
V["EX"] = [
    ("STUDYID", A, "Sponsor", "", "", "", ""),
    ("DOMAIN", A, "Sponsor", "", "", "", ""),
    ("USUBJID", D, "Sponsor", "", "MT.USUBJID", "", ""),
    ("EXSEQ", D, "Sponsor", "", "MT.SEQ", "", ""),
    ("EXTRT", D, "Sponsor", "", "MT.EXTRT", "COM.BLIND", s("RAND", "TRTCODE")),
    ("EXCAT", A, "Sponsor", "", "", "", "STUDY VACCINE"),
    ("EXDOSE", C, "Investigator", "", "", "", s("EX", "EXDOSE")),
    ("EXDOSU", C, "Investigator", "UNIT", "", "", s("EX", "EXDOSEU")),
    ("EXDOSFRM", A, "Sponsor", "FRM", "", "", "INJECTION"),
    ("EXDOSFRQ", A, "Sponsor", "FREQ", "", "", "ONCE"),
    ("EXROUTE", C, "Investigator", "ROUTE", "MT.UPPER", "", s("EX", "EXROUTE")),
    ("EXLOT", C, "Investigator", "", "", "", s("EX", "LOTNUM")),
    ("EXLOC", C, "Investigator", "LOC", "MT.EXLOC", "", s("EX", "EXLOC")),
    ("EXLAT", C, "Investigator", "LAT", "MT.UPPER", "", s("EX", "EXSIDE")),
    ("EPOCH", D, "Sponsor", "EPOCH", "MT.EPOCH", "", ""),
    ("EXSTDTC", C, "Investigator", "", "MT.ISODT", "", "EX.EXDAT + EX.EXTIM"),
    ("EXENDTC", C, "Investigator", "", "MT.ISODT", "", "Same as EXSTDTC"),
    ("EXSTDY", D, "Sponsor", "", "MT.DY", "", ""),
    ("EXENDY", D, "Sponsor", "", "MT.DY", "", ""),
]
V["SV"] = [
    ("STUDYID", A, "Sponsor", "", "", "", ""),
    ("DOMAIN", A, "Sponsor", "", "", "", ""),
    ("USUBJID", D, "Sponsor", "", "MT.USUBJID", "", ""),
    ("VISITNUM", D, "Sponsor", "", "MT.VISIT", "", s("VISIT", "VISIT")),
    ("VISIT", D, "Sponsor", "", "MT.VISIT", "", s("VISIT", "VISIT")),
    ("SVPRESP", A, "Sponsor", "NY", "", "", "Y"),
    ("SVOCCUR", A, "Sponsor", "NY", "", "", "Y"),
    ("VISITDY", D, "Sponsor", "", "MT.VISIT", "", "Planned day from TV"),
    ("SVSTDTC", C, "Investigator", "", "MT.ISODT", "", s("VISIT", "VISDAT")),
    ("SVENDTC", C, "Investigator", "", "MT.ISODT", "", "Same as SVSTDTC"),
    ("SVSTDY", D, "Sponsor", "", "MT.DY", "", ""),
    ("SVENDY", D, "Sponsor", "", "MT.DY", "", ""),
]
V["IE"] = [
    ("STUDYID", A, "Sponsor", "", "", "", ""),
    ("DOMAIN", A, "Sponsor", "", "", "", ""),
    ("USUBJID", D, "Sponsor", "", "MT.USUBJID", "", ""),
    ("IESEQ", D, "Sponsor", "", "MT.SEQ", "", ""),
    ("IETESTCD", D, "Sponsor", "", "MT.IECRIT", "", s("IE", "IETERM")),
    ("IETEST", D, "Sponsor", "", "MT.IECRIT", "", "From TI"),
    ("IECAT", D, "Sponsor", "", "MT.IECRIT", "", "INCLUSION or EXCLUSION"),
    ("IEORRES", D, "Sponsor", "NY", "MT.IECRIT", "", "Y = exclusion met, N = inclusion not met"),
    ("IESTRESC", D, "Sponsor", "NY", "MT.IECRIT", "", "= IEORRES"),
    ("VISITNUM", A, "Sponsor", "", "", "", "1"),
    ("VISIT", A, "Sponsor", "", "", "", "SCREENING"),
    ("EPOCH", A, "Sponsor", "EPOCH", "", "", "SCREENING"),
    ("IEDTC", C, "Investigator", "", "MT.ISODT", "", s("IE", "IEDAT")),
    ("IEDY", D, "Sponsor", "", "MT.DY", "", ""),
]
V["VS"] = [
    ("STUDYID", A, "Sponsor", "", "", "", ""),
    ("DOMAIN", A, "Sponsor", "", "", "", ""),
    ("USUBJID", D, "Sponsor", "", "MT.USUBJID", "", ""),
    ("VSSEQ", D, "Sponsor", "", "MT.SEQ", "", ""),
    ("VSTESTCD", D, "Sponsor", "VSTESTCD", "MT.VSTEST", "", "SYSBP, DIABP, PULSE, TEMP"),
    ("VSTEST", D, "Sponsor", "VSTEST", "MT.VSTEST", "", ""),
    ("VSCAT", D, "Sponsor", "", "MT.VSCAT", "COM.FLAT", "REACTOGENICITY for diary temperature"),
    ("VSSCAT", D, "Sponsor", "", "MT.VSCAT", "COM.FLAT", "SYSTEMIC for diary temperature"),
    ("VSORRES", C, "Investigator", "", "", "", "VITALS.SYSBP/DIABP/PULSE/TEMP, DIARY.TEMP"),
    ("VSORRESU", C, "Investigator", "VSRESU", "MT.VSUNIT", "", "VITALS.TEMPU, DIARY.TEMPU"),
    ("VSSTRESC", D, "Sponsor", "", "MT.TEMPC", "COM.TEMP", "= VSSTRESN as character"),
    ("VSSTRESN", D, "Sponsor", "", "MT.TEMPC", "COM.TEMP", ""),
    ("VSSTRESU", D, "Sponsor", "VSRESU", "MT.TEMPC", "COM.TEMP", ""),
    ("VSLOBXFL", D, "Sponsor", "NY", "MT.LOBXFL", "", ""),
    ("VISITNUM", D, "Sponsor", "", "MT.VISIT", "", "Null for diary"),
    ("VISIT", D, "Sponsor", "", "MT.VISIT", "", "Null for diary"),
    ("EPOCH", D, "Sponsor", "EPOCH", "MT.EPOCH", "", ""),
    ("VSDTC", C, "Investigator", "", "MT.ISODT", "", "VITALS.VSDAT, DIARY.DIARYDAT"),
    ("VSDY", D, "Sponsor", "", "MT.DY", "", ""),
    ("VSTPT", D, "Sponsor", "", "MT.DIARYTPT", "", "Diary rows only"),
    ("VSTPTNUM", D, "Sponsor", "", "MT.DIARYTPT", "", "Diary day, 1 to 7"),
    ("VSTPTREF", A, "Sponsor", "", "", "", "VACCINATION"),
    ("VSRFTDTC", D, "Sponsor", "", "MT.RFDOSE", "", "Dose date/time"),
]
V["CE"] = [
    ("STUDYID", A, "Sponsor", "", "", "", ""),
    ("DOMAIN", A, "Sponsor", "", "", "", ""),
    ("USUBJID", D, "Sponsor", "", "MT.USUBJID", "", ""),
    ("CESEQ", D, "Sponsor", "", "MT.SEQ", "", ""),
    ("CEGRPID", D, "Sponsor", "", "MT.REACT", "COM.FLAT", "VACCINATION 1-<reaction>"),
    ("CETERM", D, "Sponsor", "", "MT.REACT", "", "Reaction name as on the diary"),
    ("CEDECOD", D, "Sponsor", "", "MT.REACT", "", ""),
    ("CECAT", A, "Sponsor", "", "", "", "REACTOGENICITY"),
    ("CESCAT", D, "Sponsor", "", "MT.REACT", "", "LOCAL or SYSTEMIC"),
    ("CEPRESP", A, "Sponsor", "NY", "", "", "Y"),
    ("CEOCCUR", D, "Sponsor", "NY", "MT.CEOCCUR", "COM.GRADE", ""),
    ("CESEV", D, "Sponsor", "AESEV", "MT.CEOCCUR", "COM.GRADE", "Severity-scale reactions only"),
    ("CETOXGR", D, "Sponsor", "", "MT.CEOCCUR", "COM.GRADE", "Diameter and temperature reactions only"),
    ("EPOCH", D, "Sponsor", "EPOCH", "MT.EPOCH", "", ""),
    ("CESTDTC", D, "Sponsor", "", "MT.CEOCCUR", "", "First day with the reaction"),
    ("CEENDTC", D, "Sponsor", "", "MT.CEOCCUR", "", "Last day with the reaction"),
    ("CESTDY", D, "Sponsor", "", "MT.DY", "", ""),
    ("CEENDY", D, "Sponsor", "", "MT.DY", "", ""),
    ("CETPT", A, "Sponsor", "", "", "", "30 MINUTES POST-DOSE or DAY 7"),
    ("CETPTNUM", A, "Sponsor", "", "", "", "0 or 7"),
    ("CETPTREF", A, "Sponsor", "", "", "", "VACCINATION"),
    ("CERFTDTC", D, "Sponsor", "", "MT.RFDOSE", "", "Dose date/time"),
    ("CEEVINTX", A, "Sponsor", "", "", "", "WITHIN 30 MINUTES AFTER VACCINATION or WITHIN 7 DAYS AFTER VACCINATION"),
]
V["FA"] = [
    ("STUDYID", A, "Sponsor", "", "", "", ""),
    ("DOMAIN", A, "Sponsor", "", "", "", "FA, dataset name FACE"),
    ("USUBJID", D, "Sponsor", "", "MT.USUBJID", "", ""),
    ("FASEQ", D, "Sponsor", "", "MT.SEQ", "", ""),
    ("FAGRPID", D, "Sponsor", "", "MT.REACT", "COM.FLAT", "Same as CEGRPID"),
    ("FATESTCD", D, "Sponsor", "FATESTCD", "MT.FATEST", "", "OCCUR, SEV, DIAMETER"),
    ("FATEST", D, "Sponsor", "FATEST", "MT.FATEST", "", ""),
    ("FAOBJ", D, "Sponsor", "", "MT.REACT", "", "Reaction name"),
    ("FACAT", A, "Sponsor", "", "", "", "REACTOGENICITY"),
    ("FASCAT", D, "Sponsor", "", "MT.REACT", "", "LOCAL or SYSTEMIC"),
    ("FAORRES", C, "Investigator", "", "MT.FAORRES", "", "DIARY and OBS30 reaction columns"),
    ("FAORRESU", D, "Sponsor", "UNIT", "MT.FAORRES", "", "mm for DIAMETER"),
    ("FASTRESC", D, "Sponsor", "", "MT.FAORRES", "", "= FAORRES"),
    ("FASTRESN", D, "Sponsor", "", "MT.FAORRES", "", "DIAMETER only"),
    ("FASTRESU", D, "Sponsor", "UNIT", "MT.FAORRES", "", ""),
    ("FAEVAL", A, "Sponsor", "EVAL", "", "", "STUDY SUBJECT"),
    ("EPOCH", D, "Sponsor", "EPOCH", "MT.EPOCH", "", ""),
    ("FADTC", C, "Investigator", "", "MT.ISODT", "", "DIARY.DIARYDAT, OBS30.OBSDAT"),
    ("FADY", D, "Sponsor", "", "MT.DY", "", ""),
    ("FATPT", D, "Sponsor", "", "MT.DIARYTPT", "", "DAY n, or 30 MINUTES POST-DOSE"),
    ("FATPTNUM", D, "Sponsor", "", "MT.DIARYTPT", "", "Diary day, 0 for the 30-minute check"),
    ("FATPTREF", A, "Sponsor", "", "", "", "VACCINATION"),
    ("FARFTDTC", D, "Sponsor", "", "MT.RFDOSE", "", "Dose date/time"),
]
V["AE"] = [
    ("STUDYID", A, "Sponsor", "", "", "", ""),
    ("DOMAIN", A, "Sponsor", "", "", "", ""),
    ("USUBJID", D, "Sponsor", "", "MT.USUBJID", "", ""),
    ("AESEQ", D, "Sponsor", "", "MT.SEQ", "", ""),
    ("AETERM", C, "Investigator", "", "MT.TRIM", "", s("AE", "AETERM")),
    ("AEDECOD", D, "Sponsor", "", "MT.AECODE", "COM.CODING", ""),
    ("AEBODSYS", D, "Sponsor", "", "MT.AECODE", "COM.CODING", ""),
    ("AECAT", D, "Sponsor", "", "MT.AECAT", "", "UNSOLICITED"),
    ("AESEV", C, "Investigator", "AESEV", "MT.UPPER", "", s("AE", "AESEV")),
    ("AESER", C, "Investigator", "NY", "", "", s("AE", "AESER")),
    ("AEACN", D, "Sponsor", "ACN", "MT.AEMAP", "", s("AE", "AEACN")),
    ("AEREL", D, "Investigator", "", "MT.AEMAP", "", s("AE", "AEREL")),
    ("AEOUT", D, "Investigator", "OUT", "MT.AEMAP", "", s("AE", "AEOUT")),
    ("AESHOSP", C, "Investigator", "NY", "", "", s("AE", "AESHOSP")),
    ("EPOCH", D, "Sponsor", "EPOCH", "MT.EPOCH", "", ""),
    ("AESTDTC", C, "Investigator", "", "MT.ISODT", "", s("AE", "AESTDAT")),
    ("AEENDTC", C, "Investigator", "", "MT.ISODT", "", s("AE", "AEENDAT")),
    ("AESTDY", D, "Sponsor", "", "MT.DY", "", ""),
    ("AEENDY", D, "Sponsor", "", "MT.DY", "", ""),
    ("AEENRTPT", D, "Sponsor", "", "MT.AEONGO", "", s("AE", "AEONGO")),
    ("AEENTPT", D, "Sponsor", "", "MT.AEONGO", "", "RFPENDTC when ongoing"),
]
V["IS"] = [
    ("STUDYID", A, "Sponsor", "", "", "", ""),
    ("DOMAIN", A, "Sponsor", "", "", "", ""),
    ("USUBJID", D, "Sponsor", "", "MT.USUBJID", "", ""),
    ("ISSEQ", D, "Sponsor", "", "MT.SEQ", "", ""),
    ("ISTESTCD", A, "Sponsor", "ISTESTCD", "", "COM.ISTEST", "MBFAB"),
    ("ISTEST", A, "Sponsor", "ISTEST", "", "COM.ISTEST", "Functional Microbial-induced Antibody"),
    ("ISBDAGNT", D, "Sponsor", "", "MT.STRAIN", "", "From LB_HAI.ANALYTE"),
    ("ISCAT", A, "Sponsor", "", "", "", "IMMUNOGENICITY"),
    ("ISORRES", C, "Investigator", "", "MT.ISRES", "", s("LB_HAI", "RESULT")),
    ("ISORRESU", C, "Investigator", "UNIT", "MT.ISRES", "", s("LB_HAI", "UNIT")),
    ("ISSTRESC", D, "Sponsor", "", "MT.ISRES", "", "= ISORRES"),
    ("ISSTRESN", D, "Sponsor", "", "MT.ISRES", "", "Null for < LLOQ"),
    ("ISSTRESU", D, "Sponsor", "UNIT", "MT.ISRES", "", ""),
    ("ISSTAT", D, "Sponsor", "ND", "MT.ISRES", "", "NOT DONE when the result is blank"),
    ("ISREASND", D, "Sponsor", "", "MT.ISRES", "", s("LB_HAI", "COMMENT")),
    ("ISSPEC", A, "Sponsor", "SPECTYPE", "", "", "SERUM"),
    ("ISMETHOD", A, "Sponsor", "METHOD", "", "", "HEMAGGLUTINATION INHIBITION ASSAY"),
    ("ISLOBXFL", D, "Sponsor", "NY", "MT.LOBXFL", "", ""),
    ("ISLLOQ", A, "Sponsor", "", "", "", "10"),
    ("VISITNUM", D, "Sponsor", "", "MT.VISIT", "", s("LB_HAI", "VISIT")),
    ("VISIT", D, "Sponsor", "", "MT.VISIT", "", s("LB_HAI", "VISIT")),
    ("EPOCH", D, "Sponsor", "EPOCH", "MT.EPOCH", "", ""),
    ("ISDTC", C, "Investigator", "", "MT.ISODT", "", s("LB_HAI", "COLLDAT")),
    ("ISDY", D, "Sponsor", "", "MT.DY", "", ""),
]
V["TA"] = [(n, A, "Sponsor", "", "", "", "") for n in
           ["STUDYID", "DOMAIN", "ARMCD", "ARM", "TAETORD", "ETCD", "ELEMENT", "TABRANCH", "TATRANS", "EPOCH"]]
V["TE"] = [(n, A, "Sponsor", "", "", "", "") for n in
           ["STUDYID", "DOMAIN", "ETCD", "ELEMENT", "TESTRL", "TEENRL", "TEDUR"]]
V["TV"] = [(n, A, "Sponsor", "", "", "", "") for n in
           ["STUDYID", "DOMAIN", "VISITNUM", "VISIT", "VISITDY", "ARMCD", "ARM", "TVSTRL", "TVENRL"]]
V["TI"] = [(n, A, "Sponsor", "", "", "", "") for n in
           ["STUDYID", "DOMAIN", "IETESTCD", "IETEST", "IECAT", "TIVERS"]]
V["TS"] = [(n, A, "Sponsor", "", "", "", "") for n in
           ["STUDYID", "DOMAIN", "TSSEQ", "TSPARMCD", "TSPARM", "TSVAL", "TSVALNF", "TSVALCD", "TSVCDREF", "TSVCDVER"]]
V["TS"] = [(n, o, so, c if n != "TSPARMCD" else "TSPARMCD", m, co, no) for n, o, so, c, m, co, no in V["TS"]]

NUM_INT = {"AGE", "DSSEQ", "EXSEQ", "VSSEQ", "CESEQ", "FASEQ", "AESEQ", "ISSEQ", "IESEQ", "EXDOSE", "VISITDY", "TAETORD",
           "TSSEQ", "VSTPTNUM", "CETPTNUM", "FATPTNUM"}
NUM_FLOAT = {"VISITNUM", "VSSTRESN", "FASTRESN", "ISSTRESN", "ISLLOQ"}
DT = {
    "RFSTDTC": "datetime", "RFXSTDTC": "datetime", "RFXENDTC": "datetime", "EXSTDTC": "datetime",
    "EXENDTC": "datetime", "VSRFTDTC": "datetime", "CERFTDTC": "datetime", "FARFTDTC": "datetime",
    "AESTDTC": "partialDate", "AEENDTC": "partialDate", "AEENTPT": "date",
}
ASSIGNED = {"STUDYID": "VAXF101", "COUNTRY": "USA"}

var_rows = []
for dom, vs in V.items():
    for i, (n, origin, source, cl, mt, com, note) in enumerate(vs, 1):
        g = ig.get((dom, n))
        if g:
            label, role, core, typ = g["label"], g["role"], g["core"], g["type"]
        else:
            label, role = TIMING[n]
            core, typ = "Perm", ("Num" if n in ("VISITNUM", "VISITDY") or n.endswith("TPTNUM") else "Char")
        if typ == "Char":
            dtype = DT.get(n, "date" if n.endswith("DTC") else "text")
        else:
            dtype = "float" if n in NUM_FLOAT else "integer"
        assigned = ASSIGNED.get(n, dom if n == "DOMAIN" and dom != "FA" else "FACE" if n == "DOMAIN" else "")
        if n == "DOMAIN":
            assigned = dom
        var_rows.append([i, dom, n, label, dtype, "", "Yes" if core == "Req" else "No", assigned, cl, origin, source,
                         "", mt, role, com, note])

vl = [
    ("VS", "VSORRES", "VSTESTCD IN (SYSBP, DIABP, PULSE, TEMP)", "text", "", C, "Investigator", "", "Raw value as typed"),
    ("VS", "VSORRESU", "VSTESTCD EQ TEMP", "text", "VSRESU", C, "Investigator", "", "C or F"),
    ("VS", "VSORRESU", "VSTESTCD EQ SYSBP", "text", "VSRESU", C, "Investigator", "", "mmHg"),
    ("VS", "VSORRESU", "VSTESTCD EQ PULSE", "text", "VSRESU", C, "Investigator", "", "beats/min"),
    ("VS", "VSSTRESN", "VSTESTCD EQ TEMP", "float", "", D, "Sponsor", "MT.TEMPC", "Degrees C, 1 decimal"),
    ("FA", "FAORRES", "FATESTCD EQ OCCUR", "text", "NY", D, "Sponsor", "MT.FAORRES", "Y or N"),
    ("FA", "FAORRES", "FATESTCD EQ SEV", "text", "AESEV", C, "Investigator", "MT.FAORRES", "MILD, MODERATE, SEVERE"),
    ("FA", "FAORRES", "FATESTCD EQ DIAMETER", "float", "", C, "Investigator", "MT.FAORRES", "mm"),
    ("IS", "ISORRES", "ISSTAT IS NULL", "text", "", C, "Investigator", "MT.ISRES", "Dilution, or <10"),
]
vl_rows = [[i, *r[:4], "", *r[4:7], "", r[7], r[8]] for i, r in enumerate(vl, 1)]

# codelists: (id, display name, NCI codelist code, terms, CT version)
RACE_ROWS = [r for r in csv.reader(open(RACE_CSV))]
race = {"WHITE", "ASIAN", "BLACK OR AFRICAN AMERICAN", "OTHER"}
cls = [
    ("NY", "C66742", ["N", "Y"]), ("AGEU", "C66781", ["YEARS"]), ("SEX", "C66731", ["F", "M"]),
    ("ARMNULRS", "C142179", ["SCREEN FAILURE", "ASSIGNED, NOT TREATED"]),
    ("DSCAT", "C74558", ["DISPOSITION EVENT", "PROTOCOL MILESTONE"]),
    ("DSDECOD", "C66727", ["COMPLETED", "LOST TO FOLLOW-UP", "SCREEN FAILURE", "WITHDRAWAL BY SUBJECT"]),
    ("PROTMLST", "C114118", ["INFORMED CONSENT OBTAINED", "RANDOMIZED"]),
    ("EPOCH", "C99079", ["SCREENING", "TREATMENT", "FOLLOW-UP"]),
    ("UNIT", "C71620", ["mL", "mm", "titer"]), ("VSRESU", "C66770", ["C", "F", "mmHg", "beats/min"]),
    ("FRM", "C66726", ["INJECTION"]), ("FREQ", "C71113", ["ONCE"]), ("ROUTE", "C66729", ["INTRAMUSCULAR"]),
    ("LOC", "C74456", ["DELTOID MUSCLE"]), ("LAT", "C99073", ["LEFT", "RIGHT"]), ("EVAL", "C78735", ["STUDY SUBJECT"]),
    ("VSTESTCD", "C66741", ["DIABP", "PULSE", "SYSBP", "TEMP"]),
    ("VSTEST", "C67153", ["Diastolic Blood Pressure", "Pulse Rate", "Systolic Blood Pressure", "Temperature"]),
    ("FATESTCD", "C101832", ["DIAMETER", "OCCUR", "SEV"]),
    ("FATEST", "C101833", ["Diameter", "Occurrence Indicator", "Severity/Intensity"]),
    ("ISTESTCD", "C120525", ["MBFAB"]), ("ISTEST", "C120526", ["Functional Microbial-induced Antibody"]),
    ("SPECTYPE", "C78734", ["SERUM"]), ("METHOD", "C85492", ["HEMAGGLUTINATION INHIBITION ASSAY"]),
    ("AESEV", "C66769", ["MILD", "MODERATE", "SEVERE"]),
    ("OUT", "C66768", ["RECOVERED/RESOLVED", "RECOVERING/RESOLVING"]),
    ("ACN", "C66767", ["NOT APPLICABLE"]), ("ND", "C66789", ["NOT DONE"]),
    ("TSPARMCD", "C66738", ["AGEMAX", "AGEMIN", "DOSE", "DOSFRM", "DOSFRQ", "DOSU", "INDIC", "INTMODEL", "LENGTH",
                            "NARMS", "PLANSUB", "RANDOM", "ROUTE", "SDTIGVER", "SDTMVER", "SEXPOP", "SPONSOR",
                            "SSTDTC", "SENDTC", "STRATFCT", "STYPE", "TBLIND", "TCNTRL", "TINDTP", "TITLE", "TPHASE",
                            "TRT", "TTYPE", "ACTSUB"]),
    ("TPHASE", "C66737", ["PHASE II/III TRIAL"]), ("TTYPE", "C66739", ["IMMUNOGENICITY", "SAFETY"]),
    ("TBLIND", "C66735", ["DOUBLE BLIND"]), ("TCNTRL", "C66785", ["PLACEBO"]), ("INTMODEL", "C99076", ["PARALLEL"]),
    ("STYPE", "C99077", ["INTERVENTIONAL"]), ("SEXPOP", "C66732", ["BOTH"]), ("TINDTP", "C66736", ["PREVENTION"]),
]
cl_rows = []
miss = []
for cid, code, terms in cls:
    name = heads.loc[code, "clname"] if code in heads.index else cid
    sub = ct[ct["cl"] == code]
    for i, t in enumerate(terms, 1):
        m = sub[sub["val"] == t]
        if m.empty:
            miss.append((cid, t))
        tc = m.iloc[0]["code"] if len(m) else ""
        decode = {"N": "No", "Y": "Yes"}.get(t, "") if cid == "NY" else ""
        cl_rows.append([cid, name, code, "text", CT, i, t, tc, decode])
for cid in ("RACE", "ETHNIC"):
    n = 0
    for r in RACE_ROWS:
        if r[0] != cid or (cid == "RACE" and r[6] not in race):
            continue
        n += 1
        cl_rows.append([cid, r[1], r[2], "text", CT_OLD, n, r[6], r[7], ""])
assert not miss, miss

methods = [
    ("MT.USUBJID", "Algorithm", "'VAXF101-' || SUBJECT."),
    ("MT.ISODT", "Algorithm", "Raw dates DD-MON-YYYY or ISO yyyy-mm-dd -> ISO 8601. Partial 'UN-MON-YYYY' (day unknown) -> "
     "yyyy-mm. Date and time are joined when a time is collected (EX)."),
    ("MT.UPPER", "Algorithm", "Upper case of the raw text, mapped to the codelist term."),
    ("MT.TRIM", "Algorithm", "Verbatim text, leading/trailing blanks removed. Case kept."),
    ("MT.SEQ", "Algorithm", "Sequence number within USUBJID, ordered by the domain key variables."),
    ("MT.DY", "Algorithm", "--DY = (--DTC - RFSTDTC) + 1 if --DTC >= RFSTDTC, else (--DTC - RFSTDTC), using dates only. "
     "Null when either date is missing or partial. Null for subjects never dosed."),
    ("MT.RFDOSE", "Algorithm", "Date/time of the first record in EX (DOSED = Y). Null for subjects not dosed."),
    ("MT.RFEND", "Algorithm", "DSSTDTC of the DISPOSITION EVENT record."),
    ("MT.AGE", "Algorithm", "Completed years between BRTHDTC and RFICDTC."),
    ("MT.ARM", "Algorithm", "RAND.TRTCODE A -> ARMCD VAXF, ARM 'VAXF-101 0.5 mL'; B -> PBO, 'Placebo'. ACTARMCD/ACTARM "
     "equal ARMCD/ARM for dosed subjects. Screen failures: all four null, ARMNRS = SCREEN FAILURE. Randomized, never "
     "dosed: ACTARMCD/ACTARM null, ARMNRS = ASSIGNED, NOT TREATED."),
    ("MT.EPOCH", "Algorithm", "By --DTC date vs. the dose date: before -> SCREENING; on the dose date -> TREATMENT; "
     "after -> FOLLOW-UP. Partial date -> null. Subjects never dosed and screen failures -> SCREENING. Protocol "
     "milestones -> SCREENING. Pre-dose records (screening and Day 1 sample of a dosed subject: VS, IS) -> SCREENING, also "
     "on the dose date. DS disposition events -> by their date."),
    ("MT.DSDECOD", "Algorithm", "Informed consent record -> 'INFORMED CONSENT OBTAINED', randomization -> 'RANDOMIZED' "
     "(DSCAT PROTOCOL MILESTONE). Disposition: raw DSTERM mapped to NCOMPLT (Completed, Lost to Follow-up, "
     "Screen Failure, Withdrawal by Subject)."),
    ("MT.EXTRT", "Algorithm", "RAND.TRTCODE A -> 'VAXF-101', B -> 'PLACEBO'."),
    ("MT.EXLOC", "Algorithm", "'Deltoid' -> DELTOID MUSCLE."),
    ("MT.VISIT", "Table", "Raw VISIT name -> VISITNUM: Screening 1, Day 1 2, Day 8 3, Day 22 4, Day 91 5. VISITDY from TV "
     "(Screening null)."),
    ("MT.IECRIT", "Table", "IE.IETERM -> criterion: 'Acute illness within 2 weeks of enrollment' EXCL01 (IEORRES Y); "
     "'Received another vaccine within 30 days' EXCL02 (Y); 'Positive pregnancy test' EXCL03 (Y); 'Unable to attend all "
     "scheduled visits' INCL03 (N); 'Medical condition not stable' INCL02 (N). Text from TI. One record per screen "
     "failure."),
    ("MT.VSTEST", "Table", "SYSBP / DIABP / PULSE / TEMP -> test name from the codelist. One record per column per raw row."),
    ("MT.VSCAT", "Algorithm", "Diary temperature: VSCAT = REACTOGENICITY, VSSCAT = SYSTEMIC. Visit vitals: null."),
    ("MT.VSUNIT", "Algorithm", "Unit as collected for TEMP (C or F). SYSBP mmHg, DIABP mmHg, PULSE beats/min."),
    ("MT.TEMPC", "Algorithm", "TEMP: VSSTRESN = ORRES in C; if ORRESU = F, (ORRES - 32) * 5 / 9, 1 decimal. "
     "VSSTRESU = C. Other tests: VSSTRESN = VSORRES, VSSTRESU = VSORRESU."),
    ("MT.LOBXFL", "Algorithm", "'Y' on the last non-missing record per USUBJID and test before the dose: Day 1 vitals; "
     "Day 1 HAI sample. Null for subjects never dosed."),
    ("MT.DIARYTPT", "Algorithm", "Diary row: VSTPT/FATPT = 'DAY ' || DIARYDAY, number = DIARYDAY. 30-minute observation: "
     "FATPT '30 MINUTES POST-DOSE', FATPTNUM 0."),
    ("MT.REACT", "Table", "Raw column -> reaction, category: PAIN, TENDERNESS, REDNESS_MM, SWELLING_MM, INDURATION_MM "
     "(LOCAL); FEVER, CHILLS, MALAISE, MYALGIA, HEADACHE, ARTHRALGIA, NAUSEA, VOMITING (SYSTEMIC). CEGRPID/FAGRPID = "
     "'VACCINATION 1-' || reaction."),
    ("MT.FATEST", "Algorithm", "Severity-scale reactions (pain, tenderness, chills, malaise, myalgia, headache, arthralgia, "
     "nausea, vomiting): OCCUR for every diary day, SEV only on days with a reaction. Diameter reactions (redness, "
     "swelling, induration): DIAMETER for every diary day, 0 mm when none. Fever is in VS, not FACE."),
    ("MT.FAORRES", "Algorithm", "OCCUR: Y when the raw severity is not 'None', else N. SEV: raw severity upper case. "
     "DIAMETER: raw mm. FAORRESU = mm for DIAMETER, null otherwise. FASTRESN = FAORRES for DIAMETER only."),
    ("MT.CEOCCUR", "Algorithm", "Per subject and reaction over the window (30 minutes; diary days 1-7). Grade: severity "
     "scale None/Mild/Moderate/Severe -> 0-3; diameter mm: <25 0, 25-50 1, 51-100 2, >100 3; fever (C): <38.0 0, "
     "38.0-38.4 1, 38.5-38.9 2, 39.0-40.0 3, >40.0 4. CEOCCUR = Y if the maximum grade >= 1 else N. CESEV (MILD, "
     "MODERATE, SEVERE) for severity-scale reactions, CETOXGR (1-4) for diameter and temperature reactions. CESTDTC/"
     "CEENDTC = first and last day with grade >= 1 (null when CEOCCUR = N). Window rows exist only for subjects "
     "who were dosed; days with no diary entry are not imputed."),
    ("MT.AECODE", "Table", "Verbatim (case-insensitive) -> preferred term and system organ class from a local coding "
     "table (13 preferred terms, 21 verbatim variants). See COM.CODING."),
    ("MT.AECAT", "Algorithm", "'UNSOLICITED' for every record. SAEs are flagged by AESER."),
    ("MT.AEMAP", "Table", "AEACN: 'None' -> NOT APPLICABLE (single dose already given). AEREL: upper case. AEOUT: "
     "'Recovered/Resolved' -> RECOVERED/RESOLVED, 'Recovering' -> RECOVERING/RESOLVING."),
    ("MT.AEONGO", "Algorithm", "AEONGO = 'Y': AEENDTC null, AEENRTPT = 'ONGOING', AEENTPT = RFPENDTC."),
    ("MT.STRAIN", "Algorithm", "'HAI Strain A' -> 'STRAIN A', same for B and C."),
    ("MT.ISRES", "Algorithm", "ISORRES = RESULT. Numeric titer: ISSTRESC = ISORRES, ISSTRESN numeric. '<10': ISSTRESC "
     "'<10', ISSTRESN null (LLOQ 10). Blank result: ISORRES null, ISSTAT NOT DONE, ISREASND 'INSUFFICIENT SAMPLE'. "
     "Unit 'titer'."),
]
comments = [
    ("COM.RACE", "RACE and ETHNIC use codelists C74457/C66790 from CT 2025-09-26; CT 2026-09-25 has only the "
     "as-collected codelists. Same decision as P1 and P2."),
    ("COM.ARMNRS", "SDTMIG 3.4: ARMCD/ARM/ACTARMCD/ACTARM are null when a subject is not assigned or not treated; the "
     "reason is in ARMNRS."),
    ("COM.DS", "DSDECOD for disposition events uses codelist NCOMPLT (C66727) under the umbrella id DSDECOD; "
     "protocol milestones use PROTMLST."),
    ("COM.BLIND", "The study is double-blind. EX holds the actual treatment from the randomization list; the SDTM "
     "datasets are unblinded."),
    ("COM.TEMP", "Temperature is entered in C or F per subject. VSORRES keeps the entered value and unit; VSSTRESN is "
     "always in C."),
    ("COM.FLAT", "Reactogenicity follows the flat model (CDISC Vaccine TAUG): a global record per reaction in CE, "
     "daily diary values in FA (FACE), diary temperature in VS."),
    ("COM.GRADE", "Grading follows the FDA 2007 toxicity grading scale for preventive vaccine trials. Diameters under "
     "25 mm stay in FACE but do not count as a reaction."),
    ("COM.CODING", "MedDRA is licensed and not used. AEDECOD and AEBODSYS are assigned from a small local table whose "
     "values are written like MedDRA PT and SOC names; no dictionary version applies."),
    ("COM.ISTEST", "CT has no hemagglutination-inhibition test code. ISTESTCD MBFAB (functional microbial-induced "
     "antibody) is the nearest term; the assay is given by ISMETHOD and the strain by ISBDAGNT."),
    ("COM.TD", "Trial design datasets describe the synthetic study VAXF101. Not related to any real trial."),
]
issues = [
    (1, "All", "Study design is a synthetic study based on public summary information of a comparable trial. "
     "Documented in docs/study-design.md and the README.", "Closed"),
    (2, "AE", "No MedDRA. Local coding table instead (COM.CODING).", "Closed"),
    (3, "IS", "ISTESTCD from the nearest CT term (COM.ISTEST).", "Closed"),
    (4, "DM", "RACE/ETHNIC codelists missing from CT 2026-09-25. Used CT 2025-09-26 (COM.RACE).", "Closed"),
    (5, "FA/CE", "FATPT*/FARFTDTC and CETPT*/CERFTDTC/CEEVINTX are not in the parsed IG variable table. Kept as in the "
     "pharmaversesdtm vaccine data; to be checked in the P21 report.", "Open"),
    (6, "DS", "DS.DSCOMMENT 'Withdrew before vaccination' is not carried to SDTM (no SUPPDS). DSTERM keeps "
     "'Withdrawal by Subject'.", "Closed"),
    (7, "CE", "A reaction diameter under 25 mm is below the grading threshold: FACE keeps the value, CE counts "
     "no reaction (COM.GRADE).", "Closed"),
    (8, "DM", "Age group (stratification factor) is not in SDTM. Derived in ADSL from AGE.", "Closed"),
]

wb = openpyxl.Workbook()
wb.remove(wb.active)
HFONT, HFILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")


def sheet(title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    for c in ws[1]:
        c.font, c.fill = HFONT, HFILL
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"


sheet("Study", ["Attribute", "Value"], [
    ("StudyName", "VAXF101"),
    ("StudyDescription", "Synthetic randomized, double-blind, placebo-controlled study of an influenza vaccine, "
                         "SDTMIG 3.4"),
    ("ProtocolName", "VAXF101"), ("Language", "en"),
    ("Scope", "DM, DS, EX, SV, IE, VS, CE, FA (FACE), AE, IS, TA, TE, TV, TI, TS. Learning/portfolio use only; not "
              "related to any regulatory submission. Raw data are synthetic (python/make_raw.py).")], [20, 100])
sheet("Standards", ["OID", "Name", "Type", "Publishing Set", "Version", "Status"], [
    ("STD.SDTMIG", "SDTMIG", "IG", "", "3.4", "Final"),
    ("STD.CT.SDTM", "CDISC/NCI", "CT", "SDTM", "2026-09-25", "Final"),
    ("STD.CT.SDTM.2025", "CDISC/NCI", "CT", "SDTM", "2025-09-26", "Final")], [20, 14, 8, 14, 12, 10])
sheet("Datasets", ["Dataset", "Description", "Class", "Structure", "Purpose", "Key Variables", "Standard", "Comment",
                   "Developer Notes"],
      [(d, desc, cls_, st, "Tabulation", keys, "SDTMIG 3.4", com, note) for d, desc, cls_, st, keys, com, note in datasets],
      [9, 26, 16, 40, 11, 34, 11, 18, 60])
sheet("Variables", ["Order", "Dataset", "Variable", "Label", "Data Type", "Length", "Mandatory", "Assigned Value",
                    "Codelist", "Origin", "Source", "Pages", "Method", "Role", "Comment", "Developer Notes"],
      var_rows, [6, 8, 10, 34, 11, 7, 9, 13, 10, 10, 11, 10, 13, 17, 16, 50])
sheet("ValueLevel", ["Order", "Dataset", "Variable", "Where Clause", "Data Type", "Length", "Codelist", "Origin",
                     "Source", "Pages", "Method", "Developer Notes"], vl_rows, [6, 8, 9, 50, 10, 7, 10, 10, 11, 8, 14, 50])
sheet("Codelists", ["ID", "Name", "NCI Codelist Code", "Data Type", "Terminology", "Order", "Term", "NCI Term Code",
                    "Decoded Value"], cl_rows, [10, 34, 12, 9, 22, 6, 34, 12, 30])
sheet("Methods", ["ID", "Type", "Description"], methods, [14, 10, 120])
sheet("Comments", ["ID", "Description"], comments, [18, 120])
sheet("Open Issues", ["#", "Dataset", "Issue", "Status"], issues, [4, 10, 110, 8])
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)

mt_ids = {m[0] for m in methods}
com_ids = {c[0] for c in comments}
cl_ids = {r[0] for r in cl_rows}
for r in var_rows:
    assert not r[12] or r[12] in mt_ids, r
    assert not r[14] or r[14] in com_ids, r
    assert not r[8] or r[8] in cl_ids, r
for r in vl_rows:
    assert not r[10] or r[10] in mt_ids, r
for d, *_ in datasets:
    assert d in V, d
not_ig = [(r[1], r[2]) for r in var_rows if (r[1], r[2]) not in ig]
print("variables:", len(var_rows), "| value level:", len(vl_rows), "| codelist terms:", len(cl_rows))
print("not in IG table (timing):", not_ig)
